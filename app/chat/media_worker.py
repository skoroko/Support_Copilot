from __future__ import annotations

from io import BytesIO
from multiprocessing.connection import Connection


def parse_document(kind: str, data: bytes, max_pdf_pages: int, connection: Connection) -> None:
    """Parse an untrusted document inside a disposable process."""

    try:
        if kind == "pdf":
            from pypdf import PdfReader

            reader = PdfReader(BytesIO(data))
            if len(reader.pages) > max_pdf_pages:
                raise ValueError(f"PDF exceeds the {max_pdf_pages}-page limit")
            texts = [(page.extract_text() or "").strip() for page in reader.pages]
            result = "\n\n".join(text for text in texts if text).strip()
            connection.send(("ok", result, len(reader.pages)))
            return
        if kind == "docx":
            from docx import Document

            document = Document(BytesIO(data))
            blocks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        blocks.append("\t".join(cells))
            connection.send(("ok", "\n".join(blocks).strip(), 0))
            return
        raise ValueError("unsupported parser")
    except BaseException as exc:
        connection.send(("error", type(exc).__name__, 0))
    finally:
        connection.close()
