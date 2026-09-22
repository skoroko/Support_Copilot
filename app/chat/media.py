from __future__ import annotations

import base64
import asyncio
from io import BytesIO
import multiprocessing
import zipfile
from typing import Any

from fastapi import UploadFile


MAX_DOCUMENT_CHARS = 30_000
MAX_PDF_PAGES = 50
DEFAULT_MAX_MEDIA_BYTES = 10 * 1024 * 1024
MAX_IMAGE_PIXELS = 20_000_000
MAX_IMAGE_DIMENSION = 8_000
MAX_DOCX_ENTRIES = 1_000
MAX_DOCX_EXPANDED_BYTES = 25 * 1024 * 1024
PARSER_TIMEOUT_SECONDS = 10.0
SCAN_PDF_MIN_PAGES = 5
SCAN_PDF_MAX_TEXT_CHARS = 100
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class MediaProcessingError(ValueError):
    """Base error for media that cannot be converted into an LLM content part."""


class UnsupportedMediaTypeError(MediaProcessingError):
    pass


class MediaExtractionError(MediaProcessingError):
    pass


class MediaTooLargeError(MediaProcessingError):
    pass


async def media_to_part(
    media: UploadFile,
    *,
    client: Any | None = None,
    transcription_model: str | None = None,
    max_bytes: int = DEFAULT_MAX_MEDIA_BYTES,
) -> dict[str, Any]:
    """Convert one uploaded file into an OpenAI Chat Completions content part."""

    mime = (media.content_type or "").lower().strip()
    data = await _read_bounded(media, max_bytes=max_bytes)
    if not data:
        raise MediaProcessingError("Загруженный файл пуст.")

    if mime.startswith("image/"):
        await asyncio.to_thread(_validate_image, data, mime)
        encoded = base64.b64encode(data).decode("ascii")
        return {
            "type": "image_url",
            "image_url": {"url": f"data:{mime};base64,{encoded}"},
        }

    if mime.startswith("audio/") or mime == "application/ogg":
        _validate_audio_magic(data, mime)
        if client is None:
            raise MediaProcessingError("OpenAI client is required for audio transcription.")
        if not transcription_model:
            raise MediaProcessingError("TRANSCRIPTION_MODEL is required for audio transcription.")
        transcript = await whisper_transcribe(
            data,
            media.filename or "audio.ogg",
            client=client,
            model=transcription_model,
        )
        return {
            "type": "text",
            "text": f"[пользователь сказал голосом]:\n{transcript}",
        }

    if mime == "application/pdf":
        if not data.startswith(b"%PDF-"):
            raise MediaExtractionError("Содержимое файла не соответствует формату PDF.")
        extracted = await _extract_in_process("pdf", data)
        return {"type": "text", "text": f"[документ PDF]:\n{extracted[:MAX_DOCUMENT_CHARS]}"}

    if mime == DOCX_MIME or mime.endswith("wordprocessingml.document"):
        _validate_docx_archive(data)
        extracted = await _extract_in_process("docx", data)
        return {"type": "text", "text": f"[документ DOCX]:\n{extracted[:MAX_DOCUMENT_CHARS]}"}

    raise UnsupportedMediaTypeError(f"Unsupported media type: {mime or 'unknown'}")


async def _read_bounded(media: UploadFile, *, max_bytes: int) -> bytes:
    chunks: list[bytes] = []
    total = 0
    while chunk := await media.read(1024 * 1024):
        total += len(chunk)
        if total > max_bytes:
            raise MediaTooLargeError(f"Файл превышает лимит {max_bytes // (1024 * 1024)} МиБ.")
        chunks.append(chunk)
    media.size = total
    return b"".join(chunks)


def _validate_image(data: bytes, claimed_mime: str) -> None:
    from PIL import Image, UnidentifiedImageError

    expected = {
        "image/jpeg": "JPEG",
        "image/png": "PNG",
        "image/webp": "WEBP",
    }.get(claimed_mime)
    if expected is None:
        raise UnsupportedMediaTypeError("Поддерживаются только JPEG, PNG и WebP.")
    try:
        with Image.open(BytesIO(data)) as image:
            if image.format != expected:
                raise MediaExtractionError("MIME-тип изображения не соответствует содержимому.")
            width, height = image.size
            if width > MAX_IMAGE_DIMENSION or height > MAX_IMAGE_DIMENSION or width * height > MAX_IMAGE_PIXELS:
                raise MediaTooLargeError("Размеры изображения превышают безопасный лимит.")
            image.verify()
    except (UnidentifiedImageError, OSError) as exc:
        raise MediaExtractionError("Не удалось прочитать изображение.") from exc


def _validate_audio_magic(data: bytes, claimed_mime: str) -> None:
    valid = False
    if claimed_mime in {"audio/ogg", "application/ogg"}:
        valid = data.startswith(b"OggS")
    elif claimed_mime == "audio/mpeg":
        valid = data.startswith(b"ID3") or data.startswith((b"\xff\xfb", b"\xff\xf3", b"\xff\xf2"))
    elif claimed_mime in {"audio/mp4", "audio/m4a"}:
        valid = len(data) >= 12 and data[4:8] == b"ftyp"
    if not valid:
        raise MediaExtractionError("MIME-тип аудио не соответствует содержимому.")


def _validate_docx_archive(data: bytes) -> None:
    if not data.startswith(b"PK"):
        raise MediaExtractionError("Содержимое файла не соответствует формату DOCX.")
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            names = {entry.filename for entry in entries}
            if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                raise MediaExtractionError("Архив не является документом DOCX.")
            if len(entries) > MAX_DOCX_ENTRIES:
                raise MediaTooLargeError("DOCX содержит слишком много файлов.")
            expanded = sum(entry.file_size for entry in entries)
            if expanded > MAX_DOCX_EXPANDED_BYTES:
                raise MediaTooLargeError("Распакованный DOCX превышает безопасный лимит.")
    except zipfile.BadZipFile as exc:
        raise MediaExtractionError("Не удалось прочитать DOCX-документ.") from exc


async def _extract_in_process(kind: str, data: bytes) -> str:
    return await asyncio.to_thread(_extract_in_process_sync, kind, data)


def _extract_in_process_sync(kind: str, data: bytes) -> str:
    from app.chat.media_worker import parse_document

    context = multiprocessing.get_context("spawn")
    receive, send = context.Pipe(duplex=False)
    process = context.Process(target=parse_document, args=(kind, data, MAX_PDF_PAGES, send), daemon=True)
    process.start()
    send.close()
    try:
        if not receive.poll(PARSER_TIMEOUT_SECONDS):
            process.terminate()
            process.join()
            raise MediaExtractionError("Превышено время обработки документа.")
        status, payload, page_count = receive.recv()
        process.join(timeout=0.2)
    finally:
        receive.close()
        if process.is_alive():
            process.terminate()
            process.join()
    if status != "ok":
        raise MediaExtractionError(f"Не удалось прочитать {kind.upper()}-документ.")
    extracted = str(payload).strip()
    if kind == "pdf" and page_count >= SCAN_PDF_MIN_PAGES and len(extracted) < SCAN_PDF_MAX_TEXT_CHARS:
        raise MediaExtractionError("PDF похож на скан: для него требуется OCR.")
    if not extracted:
        raise MediaExtractionError(f"В {kind.upper()} не найден извлекаемый текст.")
    return extracted


async def whisper_transcribe(
    audio_bytes: bytes,
    filename: str,
    *,
    client: Any,
    model: str,
) -> str:
    """Transcribe a supported audio file without local conversion or FFmpeg."""

    audio_file = BytesIO(audio_bytes)
    audio_file.name = filename
    result = await client.audio.transcriptions.create(model=model, file=audio_file)
    text = str(getattr(result, "text", "") or "").strip()
    if not text:
        raise MediaProcessingError("Не удалось распознать речь в аудиофайле.")
    return text


def extract_pdf_text(data: bytes, *, max_pages: int = MAX_PDF_PAGES) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - dependency is declared in pyproject.toml
        raise MediaExtractionError("Для обработки PDF установите пакет pypdf.") from exc

    try:
        reader = PdfReader(BytesIO(data))
        selected_pages = reader.pages[:max_pages]
        page_texts = [(page.extract_text() or "").strip() for page in selected_pages]
    except Exception as exc:  # noqa: BLE001 - parser errors must become a user-facing validation error
        raise MediaExtractionError("Не удалось прочитать PDF-документ.") from exc

    extracted = "\n\n".join(text for text in page_texts if text).strip()
    if len(reader.pages) >= SCAN_PDF_MIN_PAGES and len(extracted) < SCAN_PDF_MAX_TEXT_CHARS:
        raise MediaExtractionError("PDF похож на скан: для него требуется OCR.")
    if not extracted:
        raise MediaExtractionError("В PDF не найден извлекаемый текст.")
    return extracted


def extract_docx_text(data: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - dependency is declared in pyproject.toml
        raise MediaExtractionError("Для обработки DOCX установите пакет python-docx.") from exc

    try:
        document = Document(BytesIO(data))
        blocks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    blocks.append("\t".join(cells))
    except Exception as exc:  # noqa: BLE001 - parser errors must become a user-facing validation error
        raise MediaExtractionError("Не удалось прочитать DOCX-документ.") from exc

    extracted = "\n".join(blocks).strip()
    if not extracted:
        raise MediaExtractionError("В DOCX не найден извлекаемый текст.")
    return extracted
