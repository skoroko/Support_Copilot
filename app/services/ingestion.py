"""Incremental, metadata-aware ingestion for the Support Copilot corpus.

The service deliberately keeps loading, metadata enrichment and indexing in one
place.  It does not send document fragments to a hosted embedding provider
until the caller explicitly confirms that transfer.
"""

from __future__ import annotations

import json
import logging
import re
from collections import Counter
from collections.abc import Iterable, Sequence
from copy import deepcopy
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path
from threading import Lock
from typing import Any
from urllib.parse import urlparse

from llama_index.core import Document, SimpleDirectoryReader
from llama_index.core.ingestion import DocstoreStrategy, IngestionPipeline
from llama_index.core.readers.base import BaseReader
from llama_index.core.schema import MetadataMode
from llama_index.core.storage.docstore import SimpleDocumentStore
from llama_index.embeddings.openai import OpenAIEmbedding
from llama_index.readers.file import (
    FlatReader,
    MarkdownReader,
)
from llama_index.vector_stores.qdrant import QdrantVectorStore
from qdrant_client import QdrantClient
from qdrant_client.models import (
    Distance,
    FieldCondition,
    Filter,
    MatchValue,
    VectorParams,
)

from app.services.chm_reader import CHMReader
from app.services.corpus import (
    EXTRACTION_VERSION,
    CorpusCatalog,
    apply_reviewed_relations,
    apply_reviewed_transcriptions,
    block,
    file_hash,
    source_identity,
    stable_id,
)
from app.services.structured_readers import (
    EvidenceSplitter,
    StructuredHTMLReader,
    StructuredPDFReader,
)

PROJECT_ROOT = Path(__file__).resolve().parents[2]
logger = logging.getLogger(__name__)
SUPPORTED_SUFFIXES = frozenset({".pdf", ".docx", ".html", ".htm", ".md", ".txt", ".chm"})

# These values are useful for citations and filtering, but adding paths,
# timestamps, author names or document versions to a chunk embedding harms
# semantic retrieval and makes otherwise identical pages look different.
EXCLUDED_EMBED_METADATA_KEYS = frozenset(
    {
        "source",
        "file_name",
        "file_path",
        "created_at",
        "last_modified",
        "author",
        "category",
        "product",
        "version",
        "page_label",
        "page",
        "total_pages",
        "chm_member",
        "source_id", "source_hash", "location_id", "extraction_version", "extraction_status",
        "language", "section_path", "evidence_blocks", "relations", "figure_regions", "page_bbox",
        "warnings", "title", "extraction_error",
    }
)
_VERSION_RE = re.compile(r"(?ix)(?:^|[._\-\s])v(?:ersion)?[._\-\s]*(\d+(?:[._\-]\d+){0,3})(?=$|[._\-\s])")
_PRODUCT_RULES: tuple[tuple[str, re.Pattern[str]], ...] = (
    ("gts_nx", re.compile(r"(?i)(?:\bgts\s+nx\b|\bgtsnx\b)")),
    ("fea_nx", re.compile(r"(?i)(?:\bfea\s+nx\b|\bfeanx\b)")),
    ("civil", re.compile(r"(?i)\b(?:midas\s+)?civil\b")),
    ("gen", re.compile(r"(?i)\b(?:midas\s+)?gen\b")),
    ("soilworks", re.compile(r"(?i)\bsoil\s*works\b")),
)
_GENERIC_MIDAS_PRODUCT_RE = re.compile(r"(?i)\bmidas\s+([a-z0-9]+(?:\s+nx)?)\b")
_PRODUCT_ALIASES = {
    "gts": "gts_nx",
    "gtsnx": "gts_nx",
    "gts_nx": "gts_nx",
    "fea": "fea_nx",
    "feanx": "fea_nx",
    "fea_nx": "fea_nx",
    "midas_civil": "civil",
    "midas_gen": "gen",
    "midas_soilworks": "soilworks",
}
# These values control the retrieval product scope (``RAG_PRODUCT_SCOPE``);
# they are not real product keys.  A document or config that resolves to one
# of them would silently match nothing under a scoped retrieval, so it is
# rejected instead of being stored as a product tag.
RESERVED_PRODUCT_TOKENS = frozenset({"all", "any", "none"})


class ExternalDataTransferRequired(PermissionError):
    """Raised before a hosted provider receives a document fragment."""


class DocumentLoadError(RuntimeError):
    """Raised when a supported source file contains no readable text."""


class PythonDocxReader(BaseReader):
    """DOCX reader backed by the already required ``python-docx`` package.

    LlamaIndex's bundled ``DocxReader`` additionally depends on ``docx2txt``.
    Using the project dependency directly keeps DOCX ingestion available in the
    standard compose image and preserves table text as searchable content.
    """

    def load_data(self, file: Path, extra_info: dict[str, Any] | None = None) -> list[Document]:
        from docx import Document as DocxDocument

        docx = DocxDocument(str(file))
        paragraphs = [paragraph.text.strip() for paragraph in docx.paragraphs if paragraph.text.strip()]
        table_rows = [
            " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            for table in docx.tables
            for row in table.rows
        ]
        text = "\n".join(paragraphs + [row for row in table_rows if row])
        return [Document(text=text, metadata=dict(extra_info or {}))]


@dataclass(frozen=True)
class IngestionConfig:
    """Settings required by :class:`IngestionService`.

    ``docstore_path`` is persisted alongside the local application state.  It
    is the state that lets a new CLI process distinguish changed source files
    from unchanged ones while Qdrant holds the corresponding chunks.
    """

    data_dir: Path = PROJECT_ROOT / "data"
    qdrant_url: str = "http://localhost:6333"
    qdrant_api_key: str | None = None
    collection_name: str = "support_copilot"
    embedding_model: str = "text-embedding-3-large"
    embedding_dimensions: int = 1024
    openai_api_key: str | None = None
    openai_base_url: str | None = None
    chunk_size: int = 1024
    chunk_overlap: int = 64
    default_product: str = "gts_nx"
    docstore_path: Path = PROJECT_ROOT / "var" / "ingestion_docstore.json"
    corpus_catalog_path: Path | None = None
    visual_transcriptions_path: Path | None = None
    exclude_dir_names: tuple[str, ...] = ()

    def __post_init__(self) -> None:
        if not self.qdrant_url.strip():
            raise ValueError("qdrant_url must not be blank")
        if not self.collection_name.strip():
            raise ValueError("collection_name must not be blank")
        if self.embedding_dimensions < 1:
            raise ValueError("embedding_dimensions must be positive")
        if self.chunk_size < 1:
            raise ValueError("chunk_size must be positive")
        if self.chunk_overlap < 0 or self.chunk_overlap >= self.chunk_size:
            raise ValueError("chunk_overlap must be non-negative and smaller than chunk_size")
        if not canonical_product(self.default_product):
            raise ValueError("default_product must not be blank")
        if canonical_product(self.default_product) in RESERVED_PRODUCT_TOKENS:
            raise ValueError("default_product must name a product key, not a scope token")

    @classmethod
    def from_settings(cls, settings: Any) -> IngestionConfig:
        """Build configuration without forcing unrelated app settings changes.

        It intentionally reuses the existing RAG names until a future settings
        migration introduces dedicated ingestion variables.
        """

        data_dir = Path(getattr(settings, "rag_data_dir", PROJECT_ROOT / "data"))
        if not data_dir.is_absolute():
            data_dir = PROJECT_ROOT / data_dir
        docstore_path = Path(
            getattr(
                settings,
                "rag_docstore_path",
                PROJECT_ROOT / "var" / "rag_eval_chunk_1024_docstore.json",
            )
        )
        if not docstore_path.is_absolute():
            docstore_path = PROJECT_ROOT / docstore_path
        secret = getattr(settings, "qdrant_api_key", None)
        qdrant_api_key = secret.get_secret_value() if secret else None
        openai_secret = getattr(settings, "openai_api_key", None)
        openai_api_key = openai_secret.get_secret_value() if openai_secret else None
        return cls(
            data_dir=data_dir,
            qdrant_url=getattr(settings, "qdrant_url", None) or "http://localhost:6333",
            qdrant_api_key=qdrant_api_key,
            collection_name=getattr(settings, "rag_collection", "rag_eval_chunk_1024"),
            embedding_model=getattr(settings, "rag_embedding_model", "text-embedding-3-large"),
            embedding_dimensions=getattr(settings, "rag_embedding_dimensions", 1024),
            openai_api_key=openai_api_key,
            openai_base_url=getattr(settings, "openai_base_url", None),
            chunk_size=getattr(settings, "rag_chunk_size", 1024),
            chunk_overlap=getattr(settings, "rag_chunk_overlap", 64),
            default_product=getattr(settings, "rag_product", "gts_nx"),
            docstore_path=docstore_path,
            exclude_dir_names=ingest_exclude_dir_names(settings),
        )


@dataclass(frozen=True)
class IngestionResult:
    """A concise, machine-friendly report for one ingest invocation."""

    source_files: int
    loaded_documents: int
    changed_documents: int
    unchanged_documents: int
    indexed_nodes: int
    failed_files: tuple[str, ...] = field(default_factory=tuple)
    extracted_locations: int = 0
    no_text_locations: int = 0
    unreadable_locations: int = 0

    @property
    def failed_count(self) -> int:
        return len(self.failed_files)

    @property
    def summary(self) -> str:
        return (
            f"{self.changed_documents} changed, {self.unchanged_documents} unchanged, "
            f"{self.failed_count} failed"
        )

    def as_dict(self) -> dict[str, Any]:
        return {
            "source_files": self.source_files,
            "loaded_documents": self.loaded_documents,
            "changed_documents": self.changed_documents,
            "unchanged_documents": self.unchanged_documents,
            "indexed_nodes": self.indexed_nodes,
            "failed_files": list(self.failed_files),
            "summary": self.summary,
            "extracted_locations": self.extracted_locations,
            "no_text_locations": self.no_text_locations,
            "unreadable_locations": self.unreadable_locations,
        }


DEFAULT_INGEST_EXCLUDE_DIRS = ("gts_nx_help_html",)


def ingest_exclude_dir_names(settings: Any) -> tuple[str, ...]:
    """Parse ``RAG_INGEST_EXCLUDE_DIRS`` (comma-separated directory names).

    The default drops the CHM-derived HTML mirror documented in
    ``docs/data_inventory.md``: its chunks duplicate the canonical CHM
    members and compete for the same evidence slots.
    """

    raw = getattr(settings, "rag_ingest_exclude_dirs", None)
    if raw is None:
        return DEFAULT_INGEST_EXCLUDE_DIRS
    if isinstance(raw, str):
        parts = raw.split(",")
    else:
        parts = [str(item) for item in raw]
    return tuple(part.strip() for part in parts if part.strip())


def supported_files(
    path_or_paths: Path | str | Iterable[Path | str],
    *,
    exclude_dir_names: frozenset[str] | set[str] = frozenset(),
) -> list[Path]:
    """Return supported files in deterministic order, skipping prior failures.

    ``exclude_dir_names`` drops whole directories by name during the recursive
    scan (e.g. a derivative mirror like the CHM-exported ``gts_nx_help_html``).
    Files passed explicitly are still returned — the caller named them.
    """

    if isinstance(path_or_paths, (str, Path)):
        candidates: Sequence[Path | str] = [path_or_paths]
    else:
        candidates = list(path_or_paths)
    result: list[Path] = []
    for candidate in candidates:
        path = Path(candidate)
        if path.is_file():
            files = [path]
        elif path.is_dir():
            files = [
                item
                for item in path.rglob("*")
                if item.is_file()
                and not (set(item.relative_to(path).parts[:-1]) & set(exclude_dir_names))
            ]
        else:
            raise FileNotFoundError(f"Ingestion path does not exist: {path}")
        result.extend(
            item
            for item in files
            if item.suffix.lower() in SUPPORTED_SUFFIXES and not item.name.endswith(".failed")
        )
    # A file can be supplied directly as well as through its parent directory.
    return sorted({item.resolve() for item in result}, key=lambda item: str(item).casefold())


def category_for_path(path: Path, *, data_root: Path) -> str:
    """Map ``data/<category>/file`` to the category, with a safe fallback."""

    try:
        relative = path.resolve().relative_to(data_root.resolve())
    except ValueError:
        return path.parent.name or "uncategorized"
    return relative.parts[0] if len(relative.parts) > 1 else "uncategorized"


def version_from_filename(file_name: str) -> str:
    """Extract explicit semantic versions such as ``manual-v2.1.pdf``."""

    match = _VERSION_RE.search(Path(file_name).stem)
    return match.group(1).replace("_", ".").replace("-", ".") if match else "unknown"


def canonical_product(value: str | None) -> str:
    """Return a stable product key suitable for metadata and filters."""

    normalized = re.sub(r"[^a-z0-9]+", "_", str(value or "").casefold()).strip("_")
    return _PRODUCT_ALIASES.get(normalized, normalized)


def infer_product(values: Iterable[object], *, default_product: str = "gts_nx") -> str:
    """Infer a MIDAS product from metadata-like values.

    Explicit product markers always win.  Files without a marker inherit the
    configured corpus product, which keeps generic chapter names usable while
    still separating documents explicitly labelled for Civil, FEA NX, etc.
    """

    # Metadata fields and directory components are independent evidence.
    # Joining them invents product names, e.g. MIDAS/PROJECTS -> "midas projects".
    searchable = [
        re.sub(r"[_\-.]+", " ", component)
        for value in values if value is not None
        for component in re.split(r"[\\/]", str(value))
    ]
    for product, pattern in _PRODUCT_RULES:
        if any(pattern.search(component) for component in searchable):
            return product
    for component in searchable:
        if match := _GENERIC_MIDAS_PRODUCT_RE.search(component):
            return canonical_product(match.group(1))
    return canonical_product(default_product) or "unknown"


def product_for_path(
    path: Path, *, default_product: str = "gts_nx", data_root: Path | None = None
) -> str:
    """Classify a source by product without depending on one exact filename."""

    if data_root is not None:
        try:
            path = path.resolve().relative_to(data_root.resolve())
        except ValueError:
            # Explicit ingestion inputs may live outside the configured corpus.
            # Their containing directories are not part of its product scope.
            path = Path(path.name)
    return infer_product(path.parts, default_product=default_product)


def _optional_author(path: Path) -> str | None:
    try:
        if path.suffix.lower() == ".docx":
            from docx import Document as DocxDocument

            author = DocxDocument(str(path)).core_properties.author
            return str(author).strip() or None
        if path.suffix.lower() == ".pdf":
            from pypdf import PdfReader

            metadata = PdfReader(str(path)).metadata
            author = metadata.author if metadata else None
            return str(author).strip() or None
    except Exception:  # noqa: BLE001 - optional author metadata cannot erase readable content
        # Source metadata must not make a readable document fail ingestion.
        return None
    return None


def enrich_metadata(
    document: Document,
    *,
    source_path: Path,
    data_root: Path,
    default_product: str = "gts_nx",
) -> Document:
    """Attach stable citation metadata while excluding it from embeddings."""

    stat = source_path.stat()
    metadata = dict(document.metadata or {})
    metadata.update(
        {
            "source": metadata.get("source") or source_path.name,
            "file_name": source_path.name,
            "file_path": str(source_path.resolve()),
            "created_at": datetime.fromtimestamp(stat.st_ctime, tz=UTC).isoformat(),
            "last_modified": datetime.fromtimestamp(stat.st_mtime, tz=UTC).isoformat(),
            "category": category_for_path(source_path, data_root=data_root),
            "product": product_for_path(
                source_path, default_product=default_product, data_root=data_root
            ),
            "version": version_from_filename(source_path.name),
            "language": metadata.get("language") or "unknown",
            "extraction_version": metadata.get("extraction_version") or EXTRACTION_VERSION,
        }
    )
    author = metadata.get("author") if "author" in metadata else _optional_author(source_path)
    if author:
        metadata["author"] = author
    document.metadata = metadata
    document.excluded_embed_metadata_keys = sorted(
        set(document.excluded_embed_metadata_keys) | EXCLUDED_EMBED_METADATA_KEYS | set(metadata)
    )
    # Retrieval code reads citation fields directly from ``node.metadata``.
    # They therefore do not need to occupy context or reduce split capacity.
    document.excluded_llm_metadata_keys = sorted(
        set(document.excluded_llm_metadata_keys) | set(metadata)
    )
    return document


def _file_extractors() -> dict[str, Any]:
    """The explicit mapping is part of the assignment's format contract."""

    html_reader = StructuredHTMLReader()
    return {
        ".pdf": StructuredPDFReader(),
        ".docx": PythonDocxReader(),
        ".html": html_reader,
        ".htm": html_reader,
        ".md": MarkdownReader(),
        ".txt": FlatReader(),
        ".chm": CHMReader(),
    }


def _preserve_pdf_page_metadata(documents: Sequence[Document], path: Path) -> None:
    """Attach 1-based pages when the pypdf-backed reader returns one document per page."""

    if path.suffix.lower() != ".pdf":
        return
    total_pages = len(documents)
    for page, document in enumerate(documents, start=1):
        metadata = dict(document.metadata or {})
        metadata.setdefault("page", page)
        metadata.setdefault("total_pages", total_pages)
        document.metadata = metadata


def _load_one_file(
    path: Path,
    *,
    data_root: Path,
    default_product: str = "gts_nx",
) -> list[Document]:
    source_id, _relative = source_identity(path, data_root)
    source_hash = file_hash(path)
    author = _optional_author(path)
    reader = SimpleDirectoryReader(
        input_files=[path],
        file_extractor=_file_extractors(),
        filename_as_id=True,
        raise_on_error=True,
        file_metadata=lambda _path: {"source_id": source_id, "source_hash": source_hash, "author": author},
    )
    documents = reader.load_data()
    _preserve_pdf_page_metadata(documents, path)
    for index, document in enumerate(documents):
        metadata = document.metadata
        metadata.setdefault("source_id", source_id)
        metadata.setdefault("source_hash", source_hash)
        metadata.setdefault("location_id", stable_id("location", source_id, index))
        metadata.setdefault("extraction_status", "text" if document.text.strip() else "no_text")
        metadata.setdefault("evidence_blocks", [block(metadata["location_id"], 0, "paragraph", document.text).as_dict()]
                            if document.text.strip() else [])
        metadata.setdefault("relations", [])
        metadata.setdefault("section_path", [])
        # Absolute-path IDs retain legacy replacement semantics. Provenance IDs above are portable.
        document.id_ = f"{path.resolve()}_part_{index}"
    return [
        enrich_metadata(
            document,
            source_path=path,
            data_root=data_root,
            default_product=default_product,
        )
        for document in documents
    ]


def _failed_path(path: Path) -> Path:
    candidate = path.with_name(f"{path.name}.failed")
    attempt = 1
    while candidate.exists():
        candidate = path.with_name(f"{path.name}.failed.{attempt}")
        attempt += 1
    return candidate


def _is_local_qdrant_url(url: str) -> bool:
    parsed = urlparse(url)
    host = (parsed.hostname or "").casefold()
    return host in {"", "localhost", "127.0.0.1", "::1", "qdrant"}


class IngestionService:
    """Load and incrementally upsert corpus files into Qdrant through LlamaIndex."""

    def __init__(
        self,
        config: IngestionConfig,
        *,
        client: QdrantClient | None = None,
        embed_model: Any | None = None,
        docstore: SimpleDocumentStore | None = None,
        requires_external_data_transfer: bool | None = None,
    ) -> None:
        self.config = config
        self._client = client
        self._embed_model = embed_model
        self._docstore = docstore or self._load_docstore(config.docstore_path)
        self._requires_external_data_transfer = (
            embed_model is None
            if requires_external_data_transfer is None
            else requires_external_data_transfer
        )
        self._pipeline: IngestionPipeline | None = None
        self._ingest_lock = Lock()

    @classmethod
    def from_settings(cls, settings: Any, **kwargs: Any) -> IngestionService:
        return cls(IngestionConfig.from_settings(settings), **kwargs)

    @property
    def docstore(self) -> SimpleDocumentStore:
        return self._docstore

    def ingest(
        self,
        path_or_paths: Path | str | Iterable[Path | str] | None = None,
        *,
        allow_external_data_transfer: bool | None = None,
    ) -> IngestionResult:
        """Load files and UPSERT changed documents only.

        A hosted embedding model is an external data transfer.  The explicit
        flag makes that decision visible at every CLI and upload call site.
        """

        allowed = bool(allow_external_data_transfer)
        if self._requires_external_data_transfer and not allowed:
            raise ExternalDataTransferRequired(
                "Refusing to send document fragments to the embedding provider. "
                "Pass allow_external_data_transfer=True after explicit approval."
            )
        if not self._requires_external_data_transfer and not _is_local_qdrant_url(self.config.qdrant_url) and not allowed:
            raise ExternalDataTransferRequired(
                "Qdrant URL is not local. Pass allow_external_data_transfer=True after explicit approval."
            )

        source_files = supported_files(
            self.config.data_dir if path_or_paths is None else path_or_paths,
            exclude_dir_names=set(self.config.exclude_dir_names),
        )
        extracted: list[Document] = []
        failures: list[str] = []
        errors: dict[str, str] = {}
        for path in source_files:
            try:
                extracted.extend(
                    _load_one_file(
                        path,
                        data_root=self.config.data_dir,
                        default_product=self.config.default_product,
                    )
                )
            except Exception as exc:
                logger.exception("Could not load corpus file %s", path)
                failures.append(str(path))
                errors[str(path)] = type(exc).__name__

        if self.config.visual_transcriptions_path:
            records = json.loads(self.config.visual_transcriptions_path.read_text(encoding="utf-8"))
            apply_reviewed_transcriptions(extracted, records, data_root=self.config.data_dir)
        documents = [document for document in extracted
                     if document.metadata.get("extraction_status") == "text"
                     and document.get_content(metadata_mode=MetadataMode.NONE).strip()]
        catalog = CorpusCatalog.from_documents(source_files, extracted, data_root=self.config.data_dir, errors=errors)
        # Catalog-confirmed outline/continuation links must reach production chunks too.
        blocks_by_location: dict[str, list[dict[str, Any]]] = {}
        for item in catalog.blocks:
            blocks_by_location.setdefault(item["location_id"], []).append(
                {key: value for key, value in item.items() if key not in {"location_id", "source_id"}}
            )
        relations_by_from: dict[str, list[dict[str, Any]]] = {}
        for relation in catalog.relations:
            relations_by_from.setdefault(relation["from_id"], []).append(relation)
        for document in extracted:
            location_id = document.metadata["location_id"]
            document.metadata["evidence_blocks"] = blocks_by_location.get(location_id, [])
            document.metadata["relations"] = [r for identity in [location_id, *(
                b["block_id"] for b in document.metadata["evidence_blocks"]
            )] for r in relations_by_from.get(identity, [])]
        incomplete = {str(d.metadata["file_path"]) for d in extracted
                      if d.metadata.get("extraction_status") == "unreadable"}
        complete_sources = {str(p.resolve()) for p in source_files if str(p) not in errors} - incomplete

        if not documents:
            with self._ingest_lock:
                if any(source in self._source_ref_doc_ids() for source in complete_sources):
                    self._index_documents([], len(source_files), tuple(failures), complete_sources=complete_sources)
                self._persist_catalog(catalog)
            return IngestionResult(
                source_files=len(source_files),
                loaded_documents=0,
                changed_documents=0,
                unchanged_documents=0,
                indexed_nodes=0,
                failed_files=tuple(failures),
                extracted_locations=len(extracted),
                no_text_locations=sum(d.metadata.get("extraction_status") == "no_text" for d in extracted),
                unreadable_locations=sum(d.metadata.get("extraction_status") == "unreadable" for d in extracted),
            )

        # File I/O and parsing can proceed concurrently. Only the shared
        # docstore/pipeline/Qdrant transaction must be serialized per service.
        with self._ingest_lock:
            result = self._index_documents(documents, len(source_files), tuple(failures), complete_sources=complete_sources)
            self._persist_catalog(catalog)
            from dataclasses import replace
            return replace(result, extracted_locations=len(extracted),
                           no_text_locations=sum(d.metadata.get("extraction_status") == "no_text" for d in extracted),
                           unreadable_locations=sum(d.metadata.get("extraction_status") == "unreadable" for d in extracted))

    def ingest_reviewed_transcriptions(
        self,
        records: Sequence[dict[str, Any]] | Path | str | None = None,
        *,
        allow_external_data_transfer: bool | None = None,
    ) -> IngestionResult:
        """Update approved visual evidence using persisted original Documents.

        This path never reparses source PDFs. It requires a complete current
        catalog and its structured original text Documents, validates every
        source hash, and upserts only pages changed by exact approved records.
        Legacy/incomplete stores must be rebuilt through :meth:`ingest` first.
        A repeat of an identical approval is a no-op; conflicting revisions
        require explicit source reingestion rather than silently replacing text.
        """
        allowed = bool(allow_external_data_transfer)
        if self._requires_external_data_transfer and not allowed:
            raise ExternalDataTransferRequired("Explicit approval is required for hosted embeddings")
        if not _is_local_qdrant_url(self.config.qdrant_url) and not allowed:
            raise ExternalDataTransferRequired("Explicit approval is required for remote Qdrant")
        if records is None:
            records = self.config.visual_transcriptions_path
        if records is None:
            raise ValueError("A reviewed transcription sidecar or records are required")
        if isinstance(records, (Path, str)):
            records = json.loads(Path(records).read_text(encoding="utf-8"))
        if not isinstance(records, (list, tuple)) or not all(isinstance(r, dict) for r in records):
            raise ValueError("Reviewed transcriptions must be a list of records")

        with self._ingest_lock:
            catalog_path = self.config.corpus_catalog_path or self.config.docstore_path.with_suffix(".corpus.json")
            if not catalog_path.is_file():
                raise ValueError("A complete structured corpus catalog is required; run full ingestion first")
            catalog = CorpusCatalog.load(catalog_path)
            originals = self._validated_review_cache(catalog)
            # Never mutate the persisted originals while validating a sidecar.
            changed: dict[str, Document] = {}
            seen_locations: set[str] = set()
            for record in records:
                if record.get("approval_status") != "approved":
                    continue
                has_locator = "visual_locator" in record
                candidates = [loc for loc in catalog.locations
                              if loc.get("source_hash") == record.get("source_hash")
                              and loc.get("page") == record.get("page")
                              and loc.get("product") == record.get("product")
                              and loc.get("version") == record.get("version")
                              and (not has_locator or loc.get("location_id") == record.get("location_id"))]
                if len(candidates) != 1 or candidates[0]["location_id"] not in originals:
                    raise ValueError("Approval requires one exact persisted source/page/product/version")
                location = candidates[0]
                lid = location["location_id"]
                for key in ("source_id", "location_id"):
                    if record.get(key) is not None and record[key] != location[key]:
                        raise ValueError("Approval contains a conflicting source/location identity")
                if record.get("relative_path") is not None and record["relative_path"] != next(
                    s["relative_path"] for s in catalog.sources if s["source_id"] == location["source_id"]
                ):
                    raise ValueError("Approval contains a conflicting relative source path")
                document = changed.get(lid) or deepcopy(originals[lid])
                if has_locator:
                    from app.services.reviewed_visual_regions import (
                        apply_reviewed_visual_region,
                    )

                    seen_locations.add(lid)
                    if apply_reviewed_visual_region(document, record, self.config.data_dir):
                        document.excluded_embed_metadata_keys = sorted(set(document.excluded_embed_metadata_keys)
                                                                       | set(document.metadata))
                        document.excluded_llm_metadata_keys = sorted(set(document.excluded_llm_metadata_keys)
                                                                     | set(document.metadata))
                        changed[lid] = document
                    continue
                images = [b for b in document.metadata["evidence_blocks"]
                          if b["kind"] == "visual_unextracted" and b.get("bbox") == record.get("bbox")
                          and b.get("attributes", {}).get("image_sha256") == record.get("image_sha256")]
                if len(images) != 1:
                    raise ValueError("Approval does not match one exact persisted image region")
                if record.get("region_id") is not None and record["region_id"] != images[0]["block_id"]:
                    raise ValueError("Approval contains a conflicting image region identity")
                existing_id = images[0].get("attributes", {}).get("approved_transcription_id")
                seen_locations.add(lid)
                if existing_id:
                    existing = next((b for b in document.metadata["evidence_blocks"]
                                     if b["block_id"] == existing_id), None)
                    attrs = existing.get("attributes", {}) if existing else {}
                    if (not existing or existing["kind"] != record.get("kind")
                            or existing["text"] != record.get("text") or existing.get("bbox") != record.get("bbox")
                            or attrs.get("method") != "reviewed_visual_transcription"
                            or attrs.get("semantic_verified") is not False
                            or attrs.get("verification_scope") != "literal_image_transcription"
                            or not all(attrs.get(k) == record.get(k) and record.get(k)
                                       for k in ("review_id", "reviewer", "reviewed_at", "image_sha256"))
                            or {"kind": "transcription_of", "to_id": images[0]["block_id"]}
                            not in existing.get("relations", [])):
                        raise ValueError("Conflicting approved transcription; explicit source rebuild required")
                    if apply_reviewed_relations(document, existing, record):
                        changed[lid] = document
                    continue
                apply_reviewed_transcriptions([document], [record])
                document.excluded_embed_metadata_keys = sorted(set(document.excluded_embed_metadata_keys)
                                                               | set(document.metadata))
                document.excluded_llm_metadata_keys = sorted(set(document.excluded_llm_metadata_keys)
                                                             | set(document.metadata))
                changed[lid] = document

            if not changed:
                return IngestionResult(source_files=len(catalog.sources), loaded_documents=len(seen_locations),
                                       changed_documents=0, unchanged_documents=len(seen_locations), indexed_nodes=0,
                                       extracted_locations=len(catalog.locations))

            # Preserve all source entries, untouched locations and their review
            # observations. Only blocks/relations rooted in affected locations
            # are replaced; cross-page incoming links retain their stable IDs.
            merged = deepcopy(catalog)
            affected = set(changed)
            old_blocks = {b["block_id"] for b in merged.blocks if b["location_id"] in affected}
            merged.blocks = [b for b in merged.blocks if b["location_id"] not in affected]
            merged.relations = [r for r in merged.relations if r["from_id"] not in affected | old_blocks]
            for location in merged.locations:
                if location["location_id"] in affected:
                    document = changed[location["location_id"]]
                    location["extraction_status"] = document.metadata["extraction_status"]
                    merged.blocks.extend({"source_id": location["source_id"], "location_id": location["location_id"], **b}
                                         for b in document.metadata["evidence_blocks"])
                    merged.relations.extend({"source_id": location["source_id"], **r}
                                            for r in document.metadata.get("relations", []))
            merged._resolve_relations()
            snapshot = deepcopy(self._docstore.to_dict())
            try:
                # An incremental page batch is not a complete source
                # replacement: never apply disappeared-page cleanup here.
                result = self._index_documents(list(changed.values()), len(catalog.sources), (), complete_sources=set())
                merged.persist(catalog_path)
            except Exception:
                self._docstore = SimpleDocumentStore.from_dict(snapshot)
                self._pipeline = None
                self._persist_docstore()
                raise
            from dataclasses import replace
            return replace(result, extracted_locations=len(merged.locations))

    def _validated_review_cache(self, catalog: CorpusCatalog) -> dict[str, Document]:
        """Reject stale source bytes, partial catalogs and lossy/legacy stores."""
        if catalog.extraction_version != EXTRACTION_VERSION:
            raise ValueError("Unsupported catalog extraction version; run full ingestion first")
        current = {source_identity(p, self.config.data_dir)[0]: (p, file_hash(p))
                   for p in supported_files(self.config.data_dir,
                                            exclude_dir_names=set(self.config.exclude_dir_names))}
        sources = {s["source_id"]: s for s in catalog.sources}
        if len(sources) != len(catalog.sources) or sources.keys() != current.keys():
            raise ValueError("Catalog must cover every current permitted source exactly once")
        if any(s.get("source_hash") != current[sid][1] or s.get("status") != "extracted"
               or s.get("relative_path") != source_identity(current[sid][0], self.config.data_dir)[1]
               or s.get("extraction_version") != EXTRACTION_VERSION for sid, s in sources.items()):
            raise ValueError("Source hash changed or extraction incomplete; run full ingestion first")
        locations = {loc["location_id"]: loc for loc in catalog.locations}
        if len(locations) != len(catalog.locations) or any(
            loc.get("source_id") not in sources
            or loc.get("source_hash") != sources[loc["source_id"]]["source_hash"]
            or loc.get("extraction_status") not in {"text", "no_text"}
            for loc in locations.values()
        ):
            raise ValueError("Catalog locations have incomplete or inconsistent source provenance")
        for sid, source in sources.items():
            units = [loc for loc in locations.values() if loc["source_id"] == sid]
            if not units or dict(Counter(loc["extraction_status"] for loc in units)) != source.get("location_statuses"):
                raise ValueError("Catalog source location coverage is incomplete")
            if source["format"] == "pdf":
                totals = {loc.get("total_pages") for loc in units}
                if (len(totals) != 1 or type(next(iter(totals))) is not int
                        or next(iter(totals)) < 1
                        or sorted(loc.get("page", 0) for loc in units) != list(range(1, next(iter(totals)) + 1))):
                    raise ValueError("Catalog must preserve every physical PDF page including blank pages")
        blocks: dict[str, list[dict[str, Any]]] = {}
        for evidence in catalog.blocks:
            if (evidence.get("location_id") not in locations
                    or evidence.get("source_id") != locations[evidence["location_id"]]["source_id"]):
                raise ValueError("Catalog block is missing its exact source location")
            blocks.setdefault(evidence["location_id"], []).append(
                {k: v for k, v in evidence.items() if k not in {"source_id", "location_id"}}
            )
        originals: dict[str, Document] = {}
        for document in self._docstore.docs.values():
            metadata = document.metadata
            lid = metadata.get("location_id")
            loc = locations.get(lid)
            if (not isinstance(document, Document) or not loc or lid in originals
                    or metadata.get("extraction_version") != EXTRACTION_VERSION
                    or any(metadata.get(k) != loc.get(k) for k in
                           ("source_id", "source_hash", "page", "chm_member", "product", "version"))
                    or Path(metadata.get("file_path", "")).resolve() != current[loc["source_id"]][0]
                    or self._docstore.get_document_hash(document.id_) != document.hash
                    or document.text != "\n\n".join(b["text"] for b in blocks.get(lid, [])
                                                   if b.get("text") and b["kind"] != "outline_heading")
                    or metadata.get("evidence_blocks") != blocks.get(lid, [])):
                raise ValueError("Docstore lacks exact structured originals matching the catalog; rebuild required")
            originals[lid] = document
        expected = {lid for lid, loc in locations.items() if loc["extraction_status"] == "text"}
        if originals.keys() != expected:
            raise ValueError("Docstore is missing original extracted text locations; rebuild required")
        return originals

    def _index_documents(
        self, documents: list[Document], source_files: int, failures: tuple[str, ...],
        *, complete_sources: set[str] | None = None,
    ) -> IngestionResult:
        """Apply one ingestion batch while holding the service's state lock."""

        changed, unchanged = self._change_counts(documents)
        # The docstore retains each indexed document's source path and ID.
        # Snapshot the old IDs before UPSERTS mutates it, and only remove IDs
        # missing from a successfully indexed replacement of that same source.
        old_ids = self._source_ref_doc_ids()
        new_ids: dict[str, set[str]] = {source: set() for source in complete_sources or ()}
        for document in documents:
            source = str(document.metadata["file_path"])
            new_ids.setdefault(source, set()).add(document.id_)

        docstore_snapshot = deepcopy(self._docstore.to_dict())
        try:
            self._ensure_collection()
            nodes = self._pipeline_or_create().run(documents=documents, show_progress=False) if documents else []
            for source, current_ids in new_ids.items():
                if complete_sources is not None and source not in complete_sources:
                    continue
                for ref_doc_id in old_ids.get(source, set()) - current_ids:
                    # Qdrant must acknowledge deletion before the docstore
                    # forgets this ID, so a failed attempt remains retryable.
                    self._client_or_create().delete(
                        collection_name=self.config.collection_name,
                        points_selector=Filter(
                            must=[FieldCondition(key="doc_id", match=MatchValue(value=ref_doc_id))]
                        ),
                        wait=True,
                    )
                    self._docstore.delete_document(ref_doc_id, raise_error=False)
            self._persist_docstore()
        except Exception:
            # UPSERTS may modify the in-memory store before embeddings or Qdrant
            # complete. Keep the last persisted IDs available for a retry.
            self._docstore = SimpleDocumentStore.from_dict(docstore_snapshot)
            self._pipeline = None
            raise
        return IngestionResult(
            source_files=source_files,
            loaded_documents=len(documents),
            changed_documents=changed,
            unchanged_documents=unchanged,
            indexed_nodes=len(nodes),
            failed_files=failures,
        )

    def close(self) -> None:
        if self._client is not None:
            self._client.close()

    def _embedding_model(self) -> Any:
        if self._embed_model is None:
            self._embed_model = OpenAIEmbedding(
                model=self.config.embedding_model,
                dimensions=self.config.embedding_dimensions,
                api_key=self.config.openai_api_key,
                api_base=self.config.openai_base_url,
            )
        return self._embed_model

    def _client_or_create(self) -> QdrantClient:
        if self._client is None:
            self._client = QdrantClient(url=self.config.qdrant_url, api_key=self.config.qdrant_api_key)
        return self._client

    def _ensure_collection(self) -> None:
        client = self._client_or_create()
        if not client.collection_exists(self.config.collection_name):
            client.create_collection(
                collection_name=self.config.collection_name,
                vectors_config=VectorParams(size=self.config.embedding_dimensions, distance=Distance.COSINE),
            )
            return
        vectors = client.get_collection(self.config.collection_name).config.params.vectors
        if isinstance(vectors, dict) or int(vectors.size) != self.config.embedding_dimensions:
            raise ValueError(
                f"Collection {self.config.collection_name!r} must use one "
                f"{self.config.embedding_dimensions}-dimension vector"
            )
        if vectors.distance != Distance.COSINE:
            raise ValueError(f"Collection {self.config.collection_name!r} must use cosine distance")

    def _pipeline_or_create(self) -> IngestionPipeline:
        if self._pipeline is None:
            vector_store = QdrantVectorStore(
                client=self._client_or_create(),
                collection_name=self.config.collection_name,
                flat_metadata=False,
            )
            self._pipeline = IngestionPipeline(
                transformations=[
                    # Citation metadata stays on the node, but must not consume
                    # the semantic chunk budget (or become embedding input).
                    EvidenceSplitter(
                        chunk_size=self.config.chunk_size,
                        chunk_overlap=self.config.chunk_overlap,
                        # Keep citation fields on produced nodes.  They remain
                        # excluded from embedding/LLM text by the per-document
                        # exclusion lists set in ``enrich_metadata``.
                    ),
                    self._embedding_model(),
                ],
                vector_store=vector_store,
                docstore=self._docstore,
                docstore_strategy=DocstoreStrategy.UPSERTS,
            )
        return self._pipeline

    def _change_counts(self, documents: Sequence[Document]) -> tuple[int, int]:
        changed = sum(
            self._docstore.get_document_hash(document.id_) != document.hash for document in documents
        )
        return changed, len(documents) - changed

    def _source_ref_doc_ids(self) -> dict[str, set[str]]:
        """Group persisted page/document IDs by their exact source path."""

        by_source: dict[str, set[str]] = {}
        for node in self._docstore.docs.values():
            source = node.metadata.get("file_path")
            ref_doc_id = node.ref_doc_id or node.id_
            if source and ref_doc_id:
                by_source.setdefault(str(source), set()).add(ref_doc_id)
        return by_source

    @staticmethod
    def _load_docstore(path: Path) -> SimpleDocumentStore:
        if path.exists():
            return SimpleDocumentStore.from_persist_path(str(path))
        return SimpleDocumentStore()

    def _persist_docstore(self) -> None:
        self.config.docstore_path.parent.mkdir(parents=True, exist_ok=True)
        self._docstore.persist(str(self.config.docstore_path))

    def _persist_catalog(self, incoming: CorpusCatalog) -> None:
        """Merge coverage for partial uploads; audit handles deletion reconciliation explicitly."""
        path = self.config.corpus_catalog_path or self.config.docstore_path.with_suffix(".corpus.json")
        if path.exists():
            old = CorpusCatalog.load(path)
            replaced = {s["source_id"] for s in incoming.sources}
            incoming.sources = [s for s in old.sources if s["source_id"] not in replaced] + incoming.sources
            incoming.locations = [s for s in old.locations if s["source_id"] not in replaced] + incoming.locations
            incoming.blocks = [s for s in old.blocks if s["source_id"] not in replaced] + incoming.blocks
            incoming.relations = [s for s in old.relations if s["source_id"] not in replaced] + incoming.relations
        incoming._resolve_relations()
        incoming.persist(path)


# The alternate name reads naturally in FastAPI composition roots and keeps the
# public contract stable for the homework description.
IngestionPipelineService = IngestionService
