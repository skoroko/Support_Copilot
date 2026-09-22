from __future__ import annotations

from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from fastapi import UploadFile
from starlette.datastructures import Headers
from PIL import Image

from app.chat.media import MediaExtractionError, MediaTooLargeError, media_to_part


def upload(filename: str, content: bytes, mime: str) -> UploadFile:
    return UploadFile(filename=filename, file=BytesIO(content), headers=Headers({"content-type": mime}))


@pytest.mark.asyncio
async def test_image_becomes_data_uri_content_part() -> None:
    data = BytesIO()
    Image.new("RGB", (2, 2), color="white").save(data, format="PNG")
    part = await media_to_part(upload("screen.png", data.getvalue(), "image/png"))
    assert part["type"] == "image_url"
    assert part["image_url"]["url"].startswith("data:image/png;base64,")


@pytest.mark.asyncio
async def test_media_rejects_body_larger_than_configured_limit() -> None:
    with pytest.raises(MediaTooLargeError):
        await media_to_part(upload("screen.png", b"x" * 9, "image/png"), max_bytes=8)


@pytest.mark.asyncio
async def test_image_rejects_spoofed_mime_type() -> None:
    with pytest.raises(MediaExtractionError):
        await media_to_part(upload("screen.png", b"not-a-png", "image/png"))


@pytest.mark.asyncio
async def test_pdf_becomes_text_content_part(monkeypatch) -> None:
    async def extract(_kind: str, _data: bytes) -> str:
        return "Инструкция по установке"

    monkeypatch.setattr("app.chat.media._extract_in_process", extract)
    part = await media_to_part(upload("guide.pdf", b"%PDF-1.7\n", "application/pdf"))
    assert part == {"type": "text", "text": "[документ PDF]:\nИнструкция по установке"}


@pytest.mark.asyncio
async def test_audio_uses_configured_transcription_model() -> None:
    create = AsyncMock(return_value=SimpleNamespace(text="Не работает лицензия"))
    client = SimpleNamespace(audio=SimpleNamespace(transcriptions=SimpleNamespace(create=create)))
    part = await media_to_part(
        upload("voice.ogg", b"OggS-test", "audio/ogg"),
        client=client,
        transcription_model="whisper-1",
    )
    assert part == {"type": "text", "text": "[пользователь сказал голосом]:\nНе работает лицензия"}
    assert create.await_args.kwargs["model"] == "whisper-1"
    assert create.await_args.kwargs["file"].name == "voice.ogg"


@pytest.mark.asyncio
async def test_docx_extracts_paragraphs_and_tables() -> None:
    from docx import Document

    document = Document()
    document.add_paragraph("Порядок установки")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Шаг"
    table.cell(0, 1).text = "Проверить лицензию"
    data = BytesIO()
    document.save(data)

    part = await media_to_part(
        upload(
            "guide.docx",
            data.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    )

    assert part["type"] == "text"
    assert part["text"].startswith("[документ DOCX]:\n")
    assert "Порядок установки" in part["text"]
    assert "Шаг\tПроверить лицензию" in part["text"]
