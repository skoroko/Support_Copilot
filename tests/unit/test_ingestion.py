from __future__ import annotations

from concurrent.futures import ThreadPoolExecutor
from concurrent.futures import TimeoutError as FutureTimeoutError
from pathlib import Path
from threading import Event

import pytest
from docx import Document as DocxDocument
from llama_index.core import Document
from llama_index.core.embeddings import MockEmbedding
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from qdrant_client import QdrantClient

from app.services.ingestion import (
    EXCLUDED_EMBED_METADATA_KEYS,
    ExternalDataTransferRequired,
    IngestionConfig,
    IngestionService,
    _file_extractors,
    _load_one_file,
    category_for_path,
    infer_product,
    product_for_path,
    supported_files,
    version_from_filename,
)
from scripts.download_data import extract_chm_to_html


def make_service(tmp_path: Path) -> IngestionService:
    return IngestionService(
        IngestionConfig(
            data_dir=tmp_path / "data",
            qdrant_url="http://qdrant:6333",
            collection_name="ingestion_test",
            embedding_dimensions=8,
            chunk_size=32,
            chunk_overlap=4,
            docstore_path=tmp_path / "docstore.json",
        ),
        client=QdrantClient(":memory:"),
        embed_model=MockEmbedding(embed_dim=8),
    )


def write_pdf(path: Path, pages: list[str]) -> None:
    """Write extractable text into a real PDF without a rendering dependency."""

    writer = PdfWriter()
    for text in pages:
        page = writer.add_blank_page(width=300, height=300)
        font = writer._add_object(
            DictionaryObject(
                {
                    NameObject("/Type"): NameObject("/Font"),
                    NameObject("/Subtype"): NameObject("/Type1"),
                    NameObject("/BaseFont"): NameObject("/Helvetica"),
                }
            )
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
        )
        content = DecodedStreamObject()
        content.set_data(f"BT /F1 12 Tf 50 250 Td ({text}) Tj ET".encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(content)
    with path.open("wb") as output:
        writer.write(output)


def indexed_points(service: IngestionService) -> list:
    points, _ = service._client.scroll(
        service.config.collection_name, limit=100, with_payload=True, with_vectors=False
    )
    return points


def test_loaders_and_metadata_cover_markdown_html_docx_and_txt(tmp_path: Path) -> None:
    data = tmp_path / "data"
    category = data / "licensing"
    category.mkdir(parents=True)
    (category / "guide-v2.1.md").write_text("# Лицензия\nКлюч активируется в личном кабинете.", encoding="utf-8")
    (category / "faq.html").write_text("<html><body><p>Срок ответа — один день.</p></body></html>", encoding="utf-8")
    (category / "notes.txt").write_text("Текстовый файл базы знаний.", encoding="utf-8")
    docx_path = category / "install.docx"
    docx = DocxDocument()
    docx.core_properties.author = "MIDAS Support"
    docx.add_paragraph("Установка завершена успешно.")
    docx.save(docx_path)

    service = make_service(tmp_path)
    result = service.ingest(data)

    assert result.source_files == 4
    assert result.failed_files == ()
    assert result.changed_documents >= 4
    assert result.indexed_nodes >= 4
    stored_documents = service.docstore.docs
    assert stored_documents
    metadata = next(node.metadata for node in stored_documents.values() if node.metadata["file_name"] == "install.docx")
    assert metadata["source"] == "install.docx"
    assert metadata["category"] == "licensing"
    assert metadata["author"] == "MIDAS Support"
    assert set(EXCLUDED_EMBED_METADATA_KEYS).issubset(
        next(node.excluded_embed_metadata_keys for node in stored_documents.values())
    )
    points, _ = service._client.scroll(
        "ingestion_test",
        limit=10,
        with_payload=True,
        with_vectors=False,
    )
    assert points
    assert all((point.payload or {}).get("file_name") for point in points)


def test_second_ingest_is_idempotent_and_persists_docstore(tmp_path: Path) -> None:
    data = tmp_path / "data" / "faq"
    data.mkdir(parents=True)
    (data / "return-v1.md").write_text("Возврат возможен в течение 14 дней.", encoding="utf-8")

    first = make_service(tmp_path)
    initial = first.ingest(tmp_path / "data")
    # Recreate the service to exercise the persisted SimpleDocumentStore.
    reopened = IngestionService(
        first.config,
        client=first._client,
        embed_model=MockEmbedding(embed_dim=8),
    )
    repeated = reopened.ingest(tmp_path / "data")

    assert initial.changed_documents == 1
    assert repeated.changed_documents == 0
    assert repeated.unchanged_documents == 1
    assert repeated.indexed_nodes == 0
    assert repeated.summary.startswith("0 changed, 1 unchanged")


def test_replacing_pdf_removes_only_disappeared_source_pages(tmp_path: Path) -> None:
    data = tmp_path / "data"
    data.mkdir()
    replaced = data / "manual.pdf"
    other = data / "other.pdf"
    write_pdf(replaced, ["CURRENT_PAGE_CONTENT", "OBSOLETE_REMOVED_PAGE"])
    write_pdf(other, ["UNRELATED_DOCUMENT_PAGE"])
    service = make_service(tmp_path)

    assert service.ingest(data).loaded_documents == 3
    old_ids = service._source_ref_doc_ids()
    assert len(old_ids[str(replaced.resolve())]) == 2
    assert len(indexed_points(service)) == 3

    write_pdf(replaced, ["CURRENT_PAGE_CONTENT_UPDATED"])
    result = service.ingest(replaced)

    assert result.failed_count == 0
    assert len(indexed_points(service)) == 2
    assert len(service._source_ref_doc_ids()[str(replaced.resolve())]) == 1
    assert service._source_ref_doc_ids()[str(other.resolve())] == old_ids[str(other.resolve())]
    payloads = [point.payload or {} for point in indexed_points(service)]
    assert not any("OBSOLETE_REMOVED_PAGE" in str(payload) for payload in payloads)
    assert any("UNRELATED_DOCUMENT_PAGE" in str(payload) for payload in payloads)

    reopened = IngestionService(
        service.config, client=service._client, embed_model=MockEmbedding(embed_dim=8)
    )
    assert reopened._source_ref_doc_ids() == service._source_ref_doc_ids()
    assert reopened.ingest(replaced).changed_documents == 0


def test_failed_stale_page_delete_keeps_replacement_retryable(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    data = tmp_path / "data"
    data.mkdir()
    source = data / "manual.pdf"
    write_pdf(source, ["CURRENT_PAGE_CONTENT", "OBSOLETE_REMOVED_PAGE"])
    service = make_service(tmp_path)
    service.ingest(source)
    old_ids = service._source_ref_doc_ids()
    original_delete = service._client.delete
    failures = 0

    def fail_once(*args: object, **kwargs: object):
        nonlocal failures
        selector = kwargs["points_selector"]
        deleting_id = selector.must[0].match.value
        if deleting_id.endswith("_part_1") and failures == 0:
            failures += 1
            raise RuntimeError("temporary Qdrant failure")
        return original_delete(*args, **kwargs)

    write_pdf(source, ["CURRENT_PAGE_CONTENT_UPDATED"])
    monkeypatch.setattr(service._client, "delete", fail_once)
    with pytest.raises(RuntimeError, match="temporary Qdrant failure"):
        service.ingest(source)
    assert service._source_ref_doc_ids() == old_ids

    retry = service.ingest(source)
    assert retry.failed_count == 0
    assert len(indexed_points(service)) == 1
    assert len(service._source_ref_doc_ids()[str(source.resolve())]) == 1


def test_concurrent_failure_cannot_erase_another_upload_from_docstore(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    data = tmp_path / "data"
    data.mkdir()
    first = data / "first.pdf"
    second = data / "second.pdf"
    write_pdf(first, ["A_CURRENT", "A_OBSOLETE"])
    service = make_service(tmp_path)
    service.ingest(first)
    write_pdf(first, ["A_REPLACEMENT"])
    write_pdf(second, ["B_CURRENT", "B_OBSOLETE_PAGE"])

    a_at_delete = Event()
    release_a = Event()
    b_loaded = Event()
    original_delete = service._client.delete
    original_load = _load_one_file

    def pause_and_fail(*args: object, **kwargs: object):
        selector = kwargs["points_selector"]
        deleting_id = selector.must[0].match.value
        if deleting_id.endswith("first.pdf_part_1"):
            a_at_delete.set()
            assert release_a.wait(timeout=10)
            raise RuntimeError("A delete failed")
        return original_delete(*args, **kwargs)

    def load_and_signal(path: Path, **kwargs: object) -> list[Document]:
        loaded = original_load(path, **kwargs)
        if path == second.resolve():
            b_loaded.set()
        return loaded

    monkeypatch.setattr(service._client, "delete", pause_and_fail)
    monkeypatch.setattr("app.services.ingestion._load_one_file", load_and_signal)

    with ThreadPoolExecutor(max_workers=2) as executor:
        a_future = executor.submit(service.ingest, first)
        assert a_at_delete.wait(timeout=10)
        b_future = executor.submit(service.ingest, second)
        assert b_loaded.wait(timeout=10)
        try:
            # B has finished parsing but cannot commit over A's snapshot.
            with pytest.raises(FutureTimeoutError):
                b_future.result(timeout=1)
        finally:
            release_a.set()
        with pytest.raises(RuntimeError, match="A delete failed"):
            a_future.result(timeout=10)
        assert b_future.result(timeout=10).loaded_documents == 2

    b_ids = service._source_ref_doc_ids()[str(second.resolve())]
    assert len(b_ids) == 2
    b_obsolete_id = next(doc_id for doc_id in b_ids if doc_id.endswith("_part_1"))
    assert b_obsolete_id in {(point.payload or {}).get("doc_id") for point in indexed_points(service)}

    write_pdf(second, ["B_REPLACEMENT"])
    service.ingest(second)
    assert b_obsolete_id not in service._source_ref_doc_ids()[str(second.resolve())]
    assert b_obsolete_id not in {(point.payload or {}).get("doc_id") for point in indexed_points(service)}


def test_failure_is_renamed_and_does_not_abort_other_files(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    data = tmp_path / "data"
    data.mkdir()
    good = data / "good.md"
    bad = data / "bad.md"
    good.write_text("Корректный текст.", encoding="utf-8")
    bad.write_text("Не будет прочитан.", encoding="utf-8")
    original = __import__("app.services.ingestion", fromlist=["_load_one_file"])._load_one_file

    def fail_bad(path: Path, **kwargs: object):
        if path.name == "bad.md":
            raise RuntimeError("fixture failure")
        return original(path, **kwargs)

    monkeypatch.setattr("app.services.ingestion._load_one_file", fail_bad)
    result = make_service(tmp_path).ingest(data)

    assert result.changed_documents == 1
    assert len(result.failed_files) == 1
    assert not bad.exists()
    assert (data / "bad.md.failed").exists()


def test_external_embedding_requires_explicit_confirmation(tmp_path: Path) -> None:
    config = IngestionConfig(data_dir=tmp_path, docstore_path=tmp_path / "docstore.json")

    with pytest.raises(ExternalDataTransferRequired):
        IngestionService(config).ingest()


def test_path_helpers_are_stable(tmp_path: Path) -> None:
    data = tmp_path / "data"
    path = data / "manuals" / "guide-v3.2.pdf"
    path.parent.mkdir(parents=True)
    path.write_text("placeholder", encoding="utf-8")

    assert category_for_path(path, data_root=data) == "manuals"
    assert version_from_filename(path.name) == "3.2"
    assert supported_files(data) == [path.resolve()]


def test_product_metadata_uses_explicit_product_markers_and_corpus_default(tmp_path: Path) -> None:
    assert product_for_path(tmp_path / "03_midas GTS NX_elements.pdf") == "gts_nx"
    assert product_for_path(tmp_path / "manuals" / "midas Civil guide.pdf") == "civil"
    assert product_for_path(tmp_path / "manuals" / "midas NFX guide.pdf") == "nfx"
    assert product_for_path(tmp_path / "manuals" / "SoilWorks manual.pdf") == "soilworks"
    assert infer_product(["generic User Manual Chapter 5.pdf"], default_product="gts_nx") == "gts_nx"


@pytest.mark.parametrize("scope", ["all", "any", "none"])
def test_ingestion_config_rejects_scope_tokens_as_default_product(tmp_path: Path, scope: str) -> None:
    with pytest.raises(ValueError, match="default_product"):
        IngestionConfig(data_dir=tmp_path, default_product=scope, docstore_path=tmp_path / "docstore.json")


def test_pdf_page_numbers_are_preserved_before_metadata_enrichment(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    pdf = tmp_path / "GTS_NX_manual.pdf"
    pdf.write_bytes(b"fixture")

    class FakeDirectoryReader:
        def __init__(self, **_kwargs: object) -> None:
            pass

        def load_data(self) -> list[Document]:
            return [
                Document(text="Первая страница"),
                Document(text="Вторая страница"),
                Document(text="Третья страница"),
            ]

    monkeypatch.setattr("app.services.ingestion.SimpleDirectoryReader", FakeDirectoryReader)

    documents = _load_one_file(pdf, data_root=tmp_path)

    assert [document.metadata["page"] for document in documents] == [1, 2, 3]
    assert all(document.metadata["total_pages"] == 3 for document in documents)
    assert all(document.metadata["product"] == "gts_nx" for document in documents)


def test_pdf_ingestion_uses_pypdf_reader() -> None:
    assert type(_file_extractors()[".pdf"]).__name__ == "PDFReader"


def test_chm_pages_can_be_materialized_as_safe_deterministic_html(tmp_path: Path) -> None:
    source = tmp_path / "manual.chm"
    source.write_bytes(b"fixture")

    class Reader:
        def load_data(self, _path: Path):
            return [
                type("Document", (), {"text": "<script>ignore()</script>Свободы", "metadata": {"chm_member": "/topic.htm"}})(),
            ]

    output = tmp_path / "data" / "gts_nx_help_html"
    files = extract_chm_to_html(source, output, reader=Reader())

    assert [item.name for item in files] == ["001-topic-htm.html"]
    rendered = files[0].read_text(encoding="utf-8")
    assert "&lt;script&gt;ignore()&lt;/script&gt;" in rendered
    assert "<script>ignore()" not in rendered
